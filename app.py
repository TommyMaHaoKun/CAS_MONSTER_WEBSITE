"""
CAS Monster Website - Flask Backend
Mirrors CAS_AUTOFILL.py functionality as a web application.
"""
import os
import io
import re
import time
import json
import secrets
import threading
from collections import defaultdict
from datetime import date as dt_date, timedelta
from typing import Optional

import requests
from flask import Flask, render_template, request, jsonify, send_from_directory
from flask_socketio import SocketIO, emit

# --------------- CAS Knowledge Base ---------------

def load_cas_knowledge():
    """Read all .docx and .xlsx files from CAS_DOCUMENTS/ and return combined text."""
    docs_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "CAS_DOCUMENTS")
    if not os.path.exists(docs_dir):
        return ""
    texts = []
    for fname in sorted(os.listdir(docs_dir)):
        if fname.startswith("~$"):
            continue
        fpath = os.path.join(docs_dir, fname)
        try:
            if fname.endswith(".docx"):
                from docx import Document as DocxDocument
                doc = DocxDocument(fpath)
                text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
                if text:
                    texts.append(f"【文档：{fname}】\n{text[:3000]}")
            elif fname.endswith(".xlsx"):
                import openpyxl
                wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        row_text = " | ".join(str(c) for c in row if c is not None)
                        if row_text.strip():
                            rows.append(row_text)
                    if rows:
                        texts.append(f"【表格：{fname} - {ws.title}】\n" + "\n".join(rows[:60]))
        except Exception as e:
            print(f"[CAS] Warning: Could not read {fname}: {e}")
    return "\n\n".join(texts)

CAS_KNOWLEDGE_BASE = load_cas_knowledge()
print(f"[CAS] Knowledge base loaded: {len(CAS_KNOWLEDGE_BASE)} chars from CAS_DOCUMENTS/")


def cas_reference_block(max_chars: int = 7000) -> str:
    if not CAS_KNOWLEDGE_BASE:
        return ""
    return (
        "\n\nOfficial school CAS reference excerpts. Follow these requirements when they are relevant:\n"
        f"{CAS_KNOWLEDGE_BASE[:max_chars]}\n"
    )

# --------------- Config ---------------
URL = "http://101.227.232.33:8001/"
# Qwen / Alibaba Cloud Model Studio. Uses the OpenAI-compatible endpoint.
LLM_BASE_URL = os.environ.get("LLM_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1").rstrip("/")
LLM_CHAT_ENDPOINT = f"{LLM_BASE_URL}/chat/completions"
LLM_FILES_ENDPOINT = f"{LLM_BASE_URL}/files"
LLM_MODEL = os.environ.get("LLM_MODEL", "qwen3.6-flash")
LLM_FALLBACK_MODEL = os.environ.get("LLM_FALLBACK_MODEL", "qwen3.7-plus")
FILE_MODEL = os.environ.get("FILE_MODEL", "qwen-long")
LLM_TEMPERATURE = float(os.environ.get("LLM_TEMPERATURE", "0.45"))
CONVERSATION_CLUB = "谈话记录(Conversation)"
# Prefer DashScope/Model Studio keys. Legacy env names stay as fallbacks only so
# existing local setups do not break immediately.
LLM_API_KEY = (os.environ.get("DASHSCOPE_API_KEY", "") or
               os.environ.get("BAILIAN_API_KEY", "") or
               os.environ.get("MOONSHOT_API_KEY", "") or
               os.environ.get("DEEPSEEK_API_KEY", ""))

# --------------- Security Config ---------------
# Max concurrent Playwright tasks to prevent resource exhaustion
MAX_CONCURRENT_TASKS = int(os.environ.get("MAX_CONCURRENT_TASKS", "3"))
_task_semaphore = threading.Semaphore(MAX_CONCURRENT_TASKS)

# Rate limiting for /api/chat (per-IP)
_RATE_LIMIT_WINDOW = 60  # seconds
_RATE_LIMIT_MAX = 15     # max requests per window
_rate_limit_store: dict[str, list[float]] = defaultdict(list)
_rate_limit_lock = threading.Lock()

PLAYWRIGHT_TIMEOUT_MS = int(os.environ.get("PLAYWRIGHT_TIMEOUT_MS", "120000"))  # 120s default

# File upload limits. Text is extracted by Qwen-Long file-extract.
ALLOWED_UPLOAD_EXT = {"txt", "md", "csv", "log", "json",
                      "docx", "pdf", "xlsx",
                      "jpg", "jpeg", "png", "webp", "bmp", "gif"}
MAX_UPLOAD_BYTES = int(os.environ.get("MAX_UPLOAD_BYTES", str(10 * 1024 * 1024)))  # 10 MB/file
MAX_UPLOAD_FILES = int(os.environ.get("MAX_UPLOAD_FILES", "5"))
MAX_EXTRACT_CHARS = int(os.environ.get("MAX_EXTRACT_CHARS", "120000"))  # per file, sent to model
QWEN_DEFAULT_CONTEXT_TOKENS = int(os.environ.get("QWEN_DEFAULT_CONTEXT_TOKENS", "1000000"))
# Qwen docs describe 1M tokens as roughly 700k Chinese characters. Use a
# character budget so chat history can use the model's default window without
# adding tokenizer dependencies.
CHAT_CONTEXT_CHAR_BUDGET = int(os.environ.get("CHAT_CONTEXT_CHAR_BUDGET", "700000"))
CHAT_MAX_MESSAGE_CHARS = int(os.environ.get("CHAT_MAX_MESSAGE_CHARS", str(CHAT_CONTEXT_CHAR_BUDGET)))

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("FLASK_SECRET_KEY", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_FILES * MAX_UPLOAD_BYTES + 2 * 1024 * 1024

# CORS: restrict to specific origins (comma-separated env var), or allow all ("*")
# when unset. "*" is the historical default and is required for the Socket.IO
# handshake to succeed when the app runs behind a proxy/cloud host (where the
# browser Origin doesn't match the server's idea of "same-origin").
_cors_origins = os.environ.get("CORS_ALLOWED_ORIGINS", "").strip()
_cors_list = [o.strip() for o in _cors_origins.split(",") if o.strip()] if _cors_origins else "*"
socketio = SocketIO(app, cors_allowed_origins=_cors_list, async_mode="threading")

# --------------- Security Helpers ---------------

def safe_str(value, max_length: int = 1000, default: str = "") -> str:
    """Ensure a value is a string with bounded length. Prevents type confusion attacks."""
    if not isinstance(value, str):
        return default
    return value.strip()[:max_length]


def sanitize_error(error: Exception) -> str:
    """Remove sensitive data (API keys, passwords) from error messages before sending to client."""
    msg = str(error)
    # Mask anything that looks like an API key or Bearer token
    msg = re.sub(r'Bearer\s+[A-Za-z0-9_\-\.]+', 'Bearer [REDACTED]', msg)
    msg = re.sub(r'sk-[A-Za-z0-9]{10,}', '[REDACTED_KEY]', msg)
    # Mask the LLM API key if it appears
    if LLM_API_KEY and len(LLM_API_KEY) > 4:
        msg = msg.replace(LLM_API_KEY, '[REDACTED_KEY]')
    return msg


def _check_rate_limit(ip: str) -> bool:
    """Returns True if request is within rate limit, False if exceeded."""
    now = time.time()
    with _rate_limit_lock:
        timestamps = _rate_limit_store[ip]
        # Remove expired entries
        _rate_limit_store[ip] = [t for t in timestamps if now - t < _RATE_LIMIT_WINDOW]
        if len(_rate_limit_store[ip]) >= _RATE_LIMIT_MAX:
            return False
        _rate_limit_store[ip].append(now)
        return True


# --------------- Security Headers ---------------

@app.after_request
def add_security_headers(response):
    """Add security headers to every response."""
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'SAMEORIGIN'
    response.headers['X-XSS-Protection'] = '1; mode=block'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    response.headers['Permissions-Policy'] = 'camera=(), microphone=(), geolocation=()'
    # Only add HSTS if behind HTTPS (check via env var)
    if os.environ.get('ENABLE_HSTS', '').lower() == 'true':
        response.headers['Strict-Transport-Security'] = 'max-age=31536000; includeSubDomains'
    return response


# --------------- Helpers ---------------

def parse_date_ymd(s: str):
    m = re.match(r"^\s*(\d{4})[/-](\d{1,2})[/-](\d{1,2})\s*$", s)
    if not m:
        raise ValueError("Date format must be YYYY/MM/DD")
    y, mo, d = map(int, m.groups())
    if not (1 <= mo <= 12 and 1 <= d <= 31):
        raise ValueError("Invalid month/day.")
    try:
        dt_date(y, mo, d)
    except ValueError as exc:
        raise ValueError("Invalid date.") from exc
    return y, mo, d


def iter_weekly_dates(start_dt: dt_date, end_dt: dt_date):
    cur = start_dt
    while cur <= end_dt:
        yield cur
        cur += timedelta(days=7)


def word_count(s: str) -> int:
    if re.search(r"[A-Za-z]", s):
        return len(re.findall(r"[A-Za-z0-9]+(?:'[A-Za-z0-9]+)?", s))
    return len(re.findall(r"\S", s))


def english_club_name(name: str) -> str:
    """Return an English-only club name for generated text/display.

    e.g. '世外信息化社(Computerization)' -> 'Computerization'. Falls back to the
    original string when no usable Latin name is present.
    """
    if not name:
        return name
    runs = re.findall(r"[A-Za-z][A-Za-z0-9 &/'\-]*", name)
    cand = max((r.strip() for r in runs), key=len, default="")
    return cand if len(cand) >= 3 else name


def parse_json_object(text: str) -> dict:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```[a-zA-Z0-9]*\n", "", cleaned)
        cleaned = re.sub(r"\n```$", "", cleaned)
        cleaned = cleaned.strip()
    try:
        return json.loads(cleaned)
    except Exception:
        m = re.search(r"\{.*\}", cleaned, re.S)
        if not m:
            raise ValueError("No JSON object found in response.")
        return json.loads(m.group(0))


# --------------- Qwen chat API ---------------

def clean_llm_messages(messages: list) -> list:
    """Drop empty chat messages before sending them to the LLM.

    Some OpenAI-compatible providers reject messages whose content is empty.
    User-visible chat history can contain empty assistant turns when a prior
    model response only produced a tool call/card.
    """
    cleaned = []
    for msg in messages or []:
        if not isinstance(msg, dict):
            continue
        role = msg.get("role")
        if role not in ("system", "user", "assistant", "tool"):
            continue
        content = msg.get("content", "")
        if content is None:
            content = ""
        if not isinstance(content, str):
            content = json.dumps(content, ensure_ascii=False)
        if not content.strip():
            continue
        new_msg = dict(msg)
        new_msg["content"] = content
        cleaned.append(new_msg)
    return cleaned


def qwen_supports_thinking(model: str) -> bool:
    model_name = (model or LLM_MODEL or "").lower()
    return model_name.startswith(("qwen3.7", "qwen3.6", "qwen3.5", "qwen3-vl", "qwen3-"))


def thinking_payload_for_model(model: str, enabled: Optional[bool]) -> Optional[dict]:
    if enabled is None:
        return None
    model_name = (model or LLM_MODEL or "").lower()
    if qwen_supports_thinking(model_name):
        return {"qwen_enable_thinking": bool(enabled)}
    if model_name.startswith("kimi-k2.7-code"):
        return {"type": "enabled"} if enabled else None
    if model_name.startswith("kimi-k2.6") or model_name.startswith("kimi-k2.5"):
        return {"type": "enabled"} if enabled else {"type": "disabled"}
    return None


def llm_chat(api_key: str, model: str, messages: list, temperature: float = 0.5,
             max_tokens: int = 600, tools: Optional[list] = None,
             tool_choice: Optional[str] = None,
             thinking_enabled: Optional[bool] = None) -> dict:
    model_name = model or LLM_MODEL
    thinking = thinking_payload_for_model(model_name, thinking_enabled)
    effective_temperature = temperature
    if not (model_name or "").lower().startswith("qwen"):
        effective_temperature = LLM_TEMPERATURE
    if thinking and thinking.get("type") == "disabled":
        effective_temperature = 0.6
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    payload = {
        "model": model_name,
        "messages": clean_llm_messages(messages),
        # Qwen accepts normal sampling temperatures; thinking is controlled per
        # request with enable_thinking.
        "temperature": effective_temperature,
        "max_tokens": max_tokens,
    }
    if thinking:
        if "qwen_enable_thinking" in thinking:
            payload["enable_thinking"] = thinking["qwen_enable_thinking"]
        else:
            payload["thinking"] = thinking
    if tools:
        payload["tools"] = tools
        if tool_choice:
            payload["tool_choice"] = tool_choice
    r = requests.post(LLM_CHAT_ENDPOINT, headers=headers, json=payload, timeout=90)
    if r.status_code != 200:
        # Surface the provider's own error message (the API key lives in the header, not
        # the response body, so including the body is safe).
        detail = ""
        try:
            j = r.json()
            err = j.get("error") if isinstance(j, dict) else None
            if isinstance(err, dict):
                detail = err.get("message") or err.get("type") or ""
            elif isinstance(j, dict):
                detail = j.get("message") or ""
            if not detail:
                detail = json.dumps(j, ensure_ascii=False)[:400]
        except Exception:
            detail = (r.text or "")[:400]
        raise RuntimeError(f"Qwen API error HTTP {r.status_code}: {sanitize_error(RuntimeError(detail))}")
    return r.json()


def llm_message_content(resp: dict) -> str:
    try:
        return (resp["choices"][0]["message"].get("content") or "").strip()
    except Exception:
        return ""


def generate_activity_record_deepseek(api_key, club_name, date_ymd, theme, c_hours, a_hours, s_hours, model=LLM_MODEL):
    return generate_activity_record_from_context_deepseek(
        api_key=api_key, club_name=club_name, date_ymd=date_ymd, theme=theme,
        c_hours=c_hours, a_hours=a_hours, s_hours=s_hours, activity_desc="",
        model=model,
    )


def generate_activity_title_deepseek(api_key, club_name, date_ymd, activity_desc, model=LLM_MODEL,
                                     thinking_enabled: bool = False):
    club_name = english_club_name(club_name)
    user_content = (
        f"Create one concise English IB CAS Activity Record title.\n"
        f"Club: {club_name}\n"
        f"Date: {date_ymd}\n"
        f"Student-provided activity description: {activity_desc}\n\n"
        f"Rules:\n"
        f"- Output ONLY the title.\n"
        f"- 4-10 words.\n"
        f"- No date, no quotes, no labels, no punctuation at the end.\n"
        f"- Make it specific to the student's description."
    )
    messages = [
        {"role": "system", "content": "Return only one short IB CAS activity title."},
        {"role": "user", "content": user_content},
    ]
    last = ""
    for _ in range(3):
        resp = llm_chat(api_key, model, messages, temperature=0.45, max_tokens=80,
                        thinking_enabled=thinking_enabled)
        raw_title = llm_message_content(resp)
        if not raw_title:
            messages.append({"role": "user", "content": "Your previous response had no final content. Output only the title."})
            continue
        title = raw_title.split("\n")[0].strip().strip('"')
        title = re.sub(r"[:：。.!?]+$", "", re.sub(r"\s+", " ", title)).strip()
        last = title
        wc = word_count(title)
        if 4 <= wc <= 10:
            return title
        messages.append({"role": "assistant", "content": title})
        messages.append({"role": "user", "content": "Revise: output only a 4-10 word title."})
    return last or "Focused CAS Activity Session"


def generate_activity_record_from_context_deepseek(api_key, club_name, date_ymd, theme, c_hours, a_hours, s_hours, activity_desc="", model=LLM_MODEL,
                                                   thinking_enabled: bool = False):
    is_conversation = club_name.strip() == CONVERSATION_CLUB
    club_name = english_club_name(club_name)
    min_words = 175 if is_conversation else 100
    word_target = "180-220" if is_conversation else "120-180"
    activity_context = f"- Student-provided activity description: {activity_desc}\n" if activity_desc else ""
    user_content = (
        f"Write an IB CAS Activity Record for the club '{club_name}'.\n"
        f"Context:\n"
        f"- Date: {date_ymd}\n"
        f"- Activity theme: {theme}\n"
        f"{activity_context}"
        f"- Hours: C={c_hours}, A={a_hours}, S={s_hours}\n\n"
        f"Hard output rules (MUST follow):\n"
        f"- Output ONLY the final record text.\n"
        f"- Do NOT add a title, labels, prefaces, explanations, word counts, or any extra lines.\n"
        f"- No bullet points, no markdown, no quotes.\n\n"
        f"Content requirements:\n"
        f"- Write ONLY in English. Do NOT output any Chinese characters. Refer to the club by its English name.\n"
        f"- Realistic high school tone.\n"
        f"- 1-2 coherent paragraphs.\n"
        f"- {word_target} words (must be >= {min_words} words).\n"
        f"- Base the concrete details on the student-provided description when it is present.\n"
        f"- The first 80% of the text must be concrete and specific.\n"
        f"- Only in the LAST 1-2 sentences, allow brief general reflection."
        f"{cas_reference_block()}"
    )
    messages = [
        {"role": "system", "content": "You write IB CAS Activity Records. Output ONLY the final prose."},
        {"role": "user", "content": user_content},
    ]
    last_text = ""
    for _ in range(3):
        resp = llm_chat(api_key, model, messages, temperature=0.55, max_tokens=520,
                        thinking_enabled=thinking_enabled)
        text = llm_message_content(resp)
        last_text = text
        if not text:
            messages.append({"role": "user", "content": "Your previous response had no final content. Output ONLY the final record text now."})
            continue
        if word_count(text) >= min_words:
            return text
        messages.append({"role": "assistant", "content": text})
        messages.append({"role": "user", "content": f"Too short. Expand to {word_target} words."})
    if not last_text:
        raise RuntimeError("Qwen returned an empty activity record. Please try again.")
    return last_text


def generate_weekly_theme_desc_deepseek(api_key, club_name, date_ymd, club_desc, periodic_desc, used_themes, used_descs=None, model=LLM_MODEL,
                                        thinking_enabled: bool = False):
    club_name = english_club_name(club_name)
    avoid = "; ".join(used_themes[-8:]) if used_themes else "none"
    if periodic_desc:
        periodic_line = f"- Periodic activity: {periodic_desc}"
        theme_guidance = (
            f"- All themes MUST be about the periodic activity: {periodic_desc}.\n"
            f"- Generate a specific aspect, session, or subtopic of this periodic activity.\n"
            f"- Keep the same overarching subject; only vary the specific focus within that subject.\n"
            f"- Avoid repetition. Do NOT reuse any themes from: {avoid}."
        )
    else:
        periodic_line = "- Periodic activity: none"
        theme_guidance = (
            f"- Avoid repetition. Do NOT reuse any themes from: {avoid}.\n"
            f"- Vary the activity focus across weeks."
        )
    user_content = (
        f"Create ONE unique Activity theme and Activity Description for the club '{club_name}'.\n"
        f"Club description: {club_desc}\n"
        f"Date: {date_ymd}\n"
        f"{periodic_line}\n\n"
        f"Return a JSON object with keys theme and description only.\n"
        f"Rules:\n"
        f"- Write ONLY in English. Do NOT output any Chinese characters. Use the club's English name.\n"
        f"- theme: 4-10 words, English, no date, no quotes.\n"
        f"- description: English, single paragraph, more than 80 words.\n"
        f"- Include at least 3 concrete details.\n"
        f"{theme_guidance}\n"
        f"- No bullet points, no markdown, no labels."
        f"{cas_reference_block()}"
    )
    messages = [
        {"role": "system", "content": "Return ONLY a valid JSON object with keys theme and description."},
        {"role": "user", "content": user_content},
    ]
    used_norm = {t.strip().lower() for t in used_themes}
    used_descs_norm = {d.strip().lower() for d in (used_descs or [])}
    last_theme, last_desc = "", ""
    for _ in range(4):
        resp = llm_chat(api_key, model, messages, temperature=0.6, max_tokens=420,
                        thinking_enabled=thinking_enabled)
        raw = llm_message_content(resp)
        if not raw:
            messages.append({"role": "user", "content": "Your previous response had no final content. Return only valid JSON with keys theme and description."})
            continue
        try:
            obj = parse_json_object(raw)
            theme = str(obj.get("theme", "")).strip()
            desc = str(obj.get("description", "")).strip()
        except Exception:
            messages.append({"role": "assistant", "content": raw})
            messages.append({"role": "user", "content": "Return only valid JSON with keys theme and description."})
            continue
        theme = re.sub(r"\s+", " ", theme)
        desc = re.sub(r"\s+", " ", desc)
        last_theme, last_desc = theme, desc
        if not theme or word_count(theme) < 4 or word_count(theme) > 10:
            messages.append({"role": "assistant", "content": raw})
            messages.append({"role": "user", "content": "Revise: theme must be 4-10 words."})
            continue
        if theme.strip().lower() in used_norm:
            messages.append({"role": "assistant", "content": raw})
            messages.append({"role": "user", "content": "Revise: theme repeats. Make it different."})
            continue
        if word_count(desc) <= 80:
            messages.append({"role": "assistant", "content": raw})
            messages.append({"role": "user", "content": "Revise: description must be more than 80 words."})
            continue
        if desc.strip().lower() in used_descs_norm:
            messages.append({"role": "assistant", "content": raw})
            messages.append({"role": "user", "content": "Revise: description repeats. Write new content."})
            continue
        return theme, desc
    return last_theme, last_desc


def generate_reflection_summary_deepseek(api_key, club_name, title, club_desc="", reflection_desc="", model=LLM_MODEL,
                                         thinking_enabled: bool = False):
    is_conversation = club_name.strip() == CONVERSATION_CLUB
    club_name = english_club_name(club_name)
    extra_context = ""
    if club_desc:
        extra_context += f"Club description: {club_desc}\n"
    if reflection_desc:
        extra_context += f"Reflection focus: {reflection_desc}\n"
    if is_conversation:
        head = (
            f"Write a concise English summary for a student's OVERALL end-of-semester IB CAS reflection.\n"
            f"Title: {title}\n\n"
            f"This summarises the student's whole CAS experience across Creativity, Activity, and Service. "
            f"It is NOT about any club; never mention a 'Conversation' or 'Conversation Club'.\n\n"
        )
    else:
        head = f"Write a concise English summary for an IB CAS reflection.\nClub: {club_name}\nTitle: {title}\n\n"
    user_content = (
        head +
        f"{extra_context}"
        f"Hard output rules:\n"
        f"- Write ONLY in English. Do NOT output any Chinese characters.\n"
        f"- Output ONLY ONE sentence.\n"
        f"- No labels, no quotes, no extra lines.\n"
        f"- About 20 words (18-22). Must end with punctuation."
        f"{cas_reference_block(3500)}"
    )
    messages = [
        {"role": "system", "content": "Return exactly one natural English sentence only."},
        {"role": "user", "content": user_content},
    ]
    last = ""
    for _ in range(4):
        resp = llm_chat(api_key, model, messages, temperature=0.45, max_tokens=100,
                        thinking_enabled=thinking_enabled)
        text = llm_message_content(resp)
        if not text:
            messages.append({"role": "user", "content": "Your previous response had no final content. Output exactly one sentence."})
            continue
        text = re.sub(r"\s+", " ", text).split("\n")[0].strip().strip('"')
        last = text
        wc = word_count(text)
        if 18 <= wc <= 22 and text.endswith(('.', '!', '?')):
            return text
        messages.append({"role": "assistant", "content": text})
        messages.append({"role": "user", "content": "Revise: 1 sentence, 18-22 words, end with punctuation."})
    return last


def generate_reflection_title_deepseek(api_key, club_name, reflection_desc, model=LLM_MODEL,
                                       thinking_enabled: bool = False):
    is_conversation = club_name.strip() == CONVERSATION_CLUB
    club_display = "Overall CAS Reflection" if is_conversation else english_club_name(club_name)
    user_content = (
        f"Create one concise English IB CAS reflection title.\n"
        f"Reflection type: {club_display}\n"
        f"Student-provided reflection/activity description: {reflection_desc}\n\n"
        f"Rules:\n"
        f"- Output ONLY the title.\n"
        f"- 4-10 words.\n"
        f"- No quotes, no labels, no punctuation at the end.\n"
        f"- Make it reflective and specific to the student's description."
    )
    messages = [
        {"role": "system", "content": "Return only one short IB CAS reflection title."},
        {"role": "user", "content": user_content},
    ]
    last = ""
    for _ in range(3):
        resp = llm_chat(api_key, model, messages, temperature=0.45, max_tokens=80,
                        thinking_enabled=thinking_enabled)
        raw_title = llm_message_content(resp)
        if not raw_title:
            messages.append({"role": "user", "content": "Your previous response had no final content. Output only the title."})
            continue
        title = raw_title.split("\n")[0].strip().strip('"')
        title = re.sub(r"[:：。.!?]+$", "", re.sub(r"\s+", " ", title)).strip()
        last = title
        wc = word_count(title)
        if 4 <= wc <= 10:
            return title
        messages.append({"role": "assistant", "content": title})
        messages.append({"role": "user", "content": "Revise: output only a 4-10 word title."})
    return last or "Growth Through CAS Experience"


def generate_reflection_content_deepseek(api_key, club_name, title, club_desc="", reflection_desc="", model=LLM_MODEL,
                                         thinking_enabled: bool = False):
    is_conversation = club_name.strip() == CONVERSATION_CLUB
    club_name = english_club_name(club_name)
    extra_context = ""
    if club_desc:
        extra_context += f"Club description: {club_desc}\n"
    if reflection_desc:
        extra_context += f"Reflection focus: {reflection_desc}\n"
    if is_conversation:
        # "谈话记录(Conversation)" is NOT a club; it is the student's holistic
        # end-of-semester CAS reflection across all three strands.
        min_words = 250
        target_phrase = "250-350"
        user_content = (
            f"Write a student's OVERALL end-of-semester IB CAS reflection in English.\n"
            f"Title: {title}\n\n"
            f"{extra_context}"
            f"CRITICAL context:\n"
            f"- This is NOT about any club. It is a holistic reflection on the student's WHOLE CAS "
            f"experience this semester across Creativity, Activity, and Service — the reflective "
            f"'conversation' a student has about their overall CAS journey.\n"
            f"- There is NO club called 'Conversation' or 'Conversation Club'. NEVER invent or mention "
            f"such a club. Do not treat 'Conversation' as an activity or organisation.\n\n"
            f"Hard output rules:\n"
            f"- Write ONLY in English. Do NOT output any Chinese characters.\n"
            f"- Output ONLY the reflection body text. No title, labels, headers, prefaces, bullets, or markdown.\n\n"
            f"Structure:\n"
            f"- At least 250 words (target 250-350).\n"
            f"- 2-4 paragraphs.\n"
            f"- Reflect across several different CAS activities/strands with concrete, specific details "
            f"and genuine personal growth.\n"
            f"- You may end with a brief forward-looking sentence."
            f"{cas_reference_block()}"
        )
    else:
        min_words = 550
        target_phrase = "600-750"
        user_content = (
            f"Write an IB CAS Activity Reflection in English.\n"
            f"Club: {club_name}\nTitle: {title}\n\n"
            f"{extra_context}"
            f"Hard output rules:\n"
            f"- Write ONLY in English. Do NOT output any Chinese characters. Use the club's English name.\n"
            f"- Output ONLY the reflection body text.\n"
            f"- No title, labels, headers, prefaces.\n"
            f"- No bullet points, no markdown.\n\n"
            f"Structure:\n"
            f"- At least 550 words (target 600-750).\n"
            f"- 4-7 paragraphs.\n"
            f"- First 70-85% must be specific with 6-10 concrete details.\n"
            f"- Only final paragraph allows general statements."
            f"{cas_reference_block()}"
        )
    messages = [
        {"role": "system", "content": "You write IB CAS reflections. Output ONLY the final body text."},
        {"role": "user", "content": user_content},
    ]
    last = ""
    for _ in range(3):
        resp = llm_chat(api_key, model, messages, temperature=0.55, max_tokens=1800,
                        thinking_enabled=thinking_enabled)
        text = llm_message_content(resp)
        last = text
        if not text:
            messages.append({"role": "user", "content": "Your previous response had no final content. Output ONLY the final reflection body text now."})
            continue
        if word_count(text) >= min_words:
            return text
        messages.append({"role": "assistant", "content": text})
        messages.append({"role": "user", "content": f"Too short. Expand to {target_phrase} words."})
    if not last:
        raise RuntimeError("Qwen returned an empty reflection. Please try again.")
    return last


# --------------- Playwright automation ---------------

def _find_iframe_src_contains(page, must_contain, timeout_ms=None):
    end = (time.time() + timeout_ms / 1000) if timeout_ms is not None else None
    while end is None or time.time() < end:
        for h in page.locator("iframe").element_handles():
            src = h.get_attribute("src") or ""
            if src and all(s in src for s in must_contain):
                return src
        time.sleep(0.12)
    return None


def _frame_locator_by_src(page, src):
    esc = src.replace("\\", "\\\\").replace('"', '\\"')
    return page.frame_locator(f'iframe[src="{esc}"]')


def int_from_text(t):
    m = re.search(r"\d+", t)
    if not m:
        raise ValueError(f"Cannot parse integer from {t!r}")
    return int(m.group())


def select_date_layui(scope, target_year, target_month, target_day):
    cal = scope.locator("#layui-laydate1")
    cal.wait_for(state="visible")

    def current_ym():
        y_txt = cal.locator(".laydate-set-ym span[lay-type='year']").inner_text()
        m_txt = cal.locator(".laydate-set-ym span[lay-type='month']").inner_text()
        return int_from_text(y_txt), int_from_text(m_txt)

    for _ in range(60):
        cy, _ = current_ym()
        if cy == target_year:
            break
        cal.locator("i.laydate-next-y" if cy < target_year else "i.laydate-prev-y").click()
        time.sleep(0.03)

    for _ in range(60):
        cy, cm = current_ym()
        if cy == target_year and cm == target_month:
            break
        cal.locator(
            "i.laydate-next-m" if (cy * 12 + cm) < (target_year * 12 + target_month) else "i.laydate-prev-m"
        ).click()
        time.sleep(0.03)

    cal.locator(f"td[lay-ymd='{target_year}-{target_month}-{target_day}']").click()


def login_and_wait_home(page, user, pw):
    page.goto(URL, wait_until="domcontentloaded")
    page.fill("input[placeholder='Please enter your login account']", user)
    page.fill("input[placeholder='Please enter your password']", pw)
    page.click("button.login-btn")
    page.wait_for_selector("text=WFLA高中综合系统")


def pick_context(page, iframe_css):
    if page.locator(iframe_css).count() > 0:
        return page.frame_locator(iframe_css)
    return page


def list_clubs_in_add_dialog(add_ctx):
    club_input = add_ctx.locator(
        "div.layui-form-item:has(label:has-text('Select a club')) "
        "div.layui-form-select input[placeholder='Please select']"
    )
    club_input.wait_for()
    club_input.click()
    options = add_ctx.locator("dd[lay-value]").filter(has_not=add_ctx.locator(".layui-select-tips"))
    options.first.wait_for()
    clubs = [t.strip() for t in options.all_inner_texts() if t.strip()]
    return [c for c in clubs if c.lower() != "please select"]


def select_club_by_text(add_ctx, club_name):
    club_input = add_ctx.locator(
        "div.layui-form-item:has(label:has-text('Select a club')) "
        "div.layui-form-select input[placeholder='Please select']"
    )
    club_input.wait_for()
    club_input.click()
    add_ctx.locator(f"dd:has-text('{club_name}')").click()


def open_records_list_ctx(page):
    page.click("text=Club Info")
    page.click("text=Activity Records")
    return pick_context(page, "iframe[src*='Stu/Cas/RecordList']")


def open_reflection_list_ctx(page):
    page.click("text=Club Info")
    page.click("text=Activity Reflection")
    src = _find_iframe_src_contains(page, ["Stu/Cas", "Reflection"])
    if src and "AddReflection" not in src:
        return _frame_locator_by_src(page, src)
    for css in ["iframe[src*='Stu/Cas/Reflection']", "iframe[src*='Stu/Cas/Reflec']"]:
        if page.locator(css).count() > 0:
            return page.frame_locator(css)
    return page


def open_add_record_ctx(record_list_ctx, page):
    btn = record_list_ctx.locator("button[data-method='add']")
    if btn.count() == 0:
        btn = record_list_ctx.locator("button:has-text('Add')")
    btn.first.click()
    add_iframe_css = "iframe[src*='/Stu/Cas/AddRecord']"
    if record_list_ctx.locator(add_iframe_css).count():
        return record_list_ctx.frame_locator(add_iframe_css)
    if page.locator(add_iframe_css).count():
        return page.frame_locator(add_iframe_css)
    return page.frame_locator(add_iframe_css)


def open_add_reflection_ctx(reflection_list_ctx, page):
    btn = reflection_list_ctx.locator("button[data-method='add']")
    if btn.count() == 0:
        btn = reflection_list_ctx.locator("button:has-text('Add')")
    btn.first.click()
    add_iframe_css = "iframe[src*='/Stu/Cas/AddReflection']"
    if reflection_list_ctx.locator(add_iframe_css).count():
        return reflection_list_ctx.frame_locator(add_iframe_css)
    if page.locator(add_iframe_css).count():
        return page.frame_locator(add_iframe_css)
    return page.frame_locator(add_iframe_css)


def fill_kindeditor_body(add_ctx, text):
    editor_iframe = add_ctx.locator("iframe.ke-edit-iframe")
    editor_iframe.first.wait_for()
    editor_frame = add_ctx.frame_locator("iframe.ke-edit-iframe")
    body = editor_frame.locator("body.ke-content")
    body.wait_for()
    body.evaluate("(el, v) => { el.innerText = v; }", text)
    body.click()
    body.press("End")


def click_learning_outcomes(add_ctx, selected):
    for name in selected:
        box = add_ctx.locator(
            f"xpath=//input[@type='checkbox' and @title='{name}']/following-sibling::div[contains(@class,'layui-form-checkbox')]"
        )
        if box.count() == 0:
            box = add_ctx.locator(f"div.layui-form-checkbox:has-text('{name}')")
        box.first.click(force=True)
        time.sleep(0.05)


# --------------- Socket.IO event handlers ---------------

def emit_log(sid, msg):
    socketio.emit("log", {"msg": msg}, to=sid)


@socketio.on("fetch_clubs")
def handle_fetch_clubs(data):
    sid = request.sid
    user = safe_str(data.get("username", ""), max_length=100)
    pw = safe_str(data.get("password", ""), max_length=200)
    if not user or not pw:
        emit("error", {"msg": "Username/Password cannot be empty."})
        return

    def task():
        acquired = _task_semaphore.acquire(timeout=5)
        if not acquired:
            emit_log(sid, "[Clubs] Server busy, too many concurrent tasks. Please try again later.")
            socketio.emit("error", {"msg": "Server busy. Please try again later."}, to=sid)
            return
        try:
            from playwright.sync_api import sync_playwright
            emit_log(sid, "[Clubs] Logging in and fetching clubs...")
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True, slow_mo=60)
                page = browser.new_page()
                page.set_default_timeout(PLAYWRIGHT_TIMEOUT_MS)
                page.set_default_navigation_timeout(PLAYWRIGHT_TIMEOUT_MS)
                login_and_wait_home(page, user, pw)
                record_list_ctx = open_records_list_ctx(page)
                add_ctx = open_add_record_ctx(record_list_ctx, page)
                clubs = list_clubs_in_add_dialog(add_ctx)
                if not clubs:
                    raise RuntimeError("No clubs found.")
                clubs_reflection = list(clubs)
                if CONVERSATION_CLUB not in clubs_reflection:
                    clubs_reflection.append(CONVERSATION_CLUB)
                browser.close()
            emit_log(sid, f"[Clubs] Fetched {len(clubs)} clubs.")
            socketio.emit("clubs_fetched", {
                "clubs_records": clubs,
                "clubs_reflection": clubs_reflection,
            }, to=sid)
        except Exception as e:
            emit_log(sid, f"[Clubs] Error: {sanitize_error(e)}")
            socketio.emit("error", {"msg": sanitize_error(e)}, to=sid)
        finally:
            _task_semaphore.release()

    threading.Thread(target=task, daemon=True).start()


@socketio.on("run_record")
def handle_run_record(data):
    sid = request.sid
    user = safe_str(data.get("username", ""), max_length=100)
    pw = safe_str(data.get("password", ""), max_length=200)
    club = safe_str(data.get("club", ""), max_length=200)
    date_str = safe_str(data.get("date", ""), max_length=20)
    theme = safe_str(data.get("theme", ""), max_length=500)
    c = safe_str(data.get("c_hours", ""), max_length=10)
    a = safe_str(data.get("a_hours", ""), max_length=10)
    s = safe_str(data.get("s_hours", ""), max_length=10)
    # Optional pre-generated (user-approved) description from the chat agent.
    approved_desc = safe_str(data.get("description", ""), max_length=5000)

    try:
        y, mo, d = parse_date_ymd(date_str)
        date_ymd = f"{y:04d}/{mo:02d}/{d:02d}"
        for name, val in [("C", c), ("A", a), ("S", s)]:
            if not re.match(r"^\d+(\.\d+)?$", val):
                raise ValueError(f"{name} hours must be a number.")
    except Exception as e:
        emit("error", {"msg": str(e)})
        return

    def task():
        acquired = _task_semaphore.acquire(timeout=5)
        if not acquired:
            emit_log(sid, "[Records] Server busy. Please try again later.")
            socketio.emit("task_done", {"task": "record"}, to=sid)
            return
        try:
            from playwright.sync_api import sync_playwright, TimeoutError as PWTimeoutError
            if approved_desc:
                desc = approved_desc
                emit_log(sid, "[Records] Using approved description. Now autofilling...")
            else:
                emit_log(sid, "[Records] Generating description via Qwen...")
                desc = generate_activity_record_deepseek(
                    api_key=LLM_API_KEY, club_name=club, date_ymd=date_ymd,
                    theme=theme, c_hours=c, a_hours=a, s_hours=s,
                )
                emit_log(sid, "[Records] Description generated. Now autofilling...")
            socketio.emit("preview_record", {"text": desc}, to=sid)

            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True, slow_mo=60)
                page = browser.new_page()
                page.set_default_timeout(PLAYWRIGHT_TIMEOUT_MS)
                page.set_default_navigation_timeout(PLAYWRIGHT_TIMEOUT_MS)
                login_and_wait_home(page, user, pw)
                record_list_ctx = open_records_list_ctx(page)
                add_ctx = open_add_record_ctx(record_list_ctx, page)

                emit_log(sid, f"[Records] Selecting club: {club}")
                select_club_by_text(add_ctx, club)

                date_input = add_ctx.locator("div.layui-form-item:has(label:has-text('Event date')) input")
                date_input.click()
                cal_scope = add_ctx if add_ctx.locator("#layui-laydate1").count() else page
                select_date_layui(cal_scope, y, mo, d)
                emit_log(sid, f"[Records] Date selected: {date_ymd}")

                add_ctx.locator("div.layui-form-item:has(label:has-text('Activity theme')) input").fill(theme)
                add_ctx.locator("input[name='CDuration']").fill(c)
                add_ctx.locator("input[name='ADuration']").fill(a)
                add_ctx.locator("input[name='SDuration']").fill(s)
                add_ctx.locator("textarea[name='Reflection']").fill(desc)
                add_ctx.locator("button[lay-filter='add']:has-text('Save')").click()
                emit_log(sid, "[Records] Save clicked.")
                time.sleep(2)
                browser.close()

            emit_log(sid, "[Records] Run finished.")
            socketio.emit("task_done", {"task": "record"}, to=sid)
        except Exception as e:
            emit_log(sid, f"[Records] Error: {sanitize_error(e)}")
            socketio.emit("task_done", {"task": "record"}, to=sid)
        finally:
            _task_semaphore.release()

    threading.Thread(target=task, daemon=True).start()


@socketio.on("run_batch")
def handle_run_batch(data):
    sid = request.sid
    user = safe_str(data.get("username", ""), max_length=100)
    pw = safe_str(data.get("password", ""), max_length=200)
    club = safe_str(data.get("club", ""), max_length=200)
    club_desc = safe_str(data.get("club_desc", ""), max_length=1000)
    weekday_label = safe_str(data.get("weekday", ""), max_length=20)
    start = safe_str(data.get("start_date", ""), max_length=20)
    end = safe_str(data.get("end_date", ""), max_length=20)
    periodic = safe_str(data.get("periodic", ""), max_length=500)
    c = safe_str(data.get("c_hours", ""), max_length=10)
    a = safe_str(data.get("a_hours", ""), max_length=10)
    s = safe_str(data.get("s_hours", ""), max_length=10)

    WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

    try:
        for name, val in [("C", c), ("A", a), ("S", s)]:
            if not re.match(r"^\d+(\.\d+)?$", val):
                raise ValueError(f"{name} hours must be a number.")
        y1, m1, d1 = parse_date_ymd(start)
        y2, m2, d2 = parse_date_ymd(end)
        start_dt = dt_date(y1, m1, d1)
        end_dt = dt_date(y2, m2, d2)
        if weekday_label not in WEEKDAYS:
            raise ValueError("Invalid weekday.")
        weekday_idx = WEEKDAYS.index(weekday_label)
        if start_dt.weekday() != weekday_idx or end_dt.weekday() != weekday_idx:
            raise ValueError(f"Start and end dates must both be {weekday_label}.")
        if start_dt > end_dt:
            raise ValueError("End date must be after start date.")
        dates = list(iter_weekly_dates(start_dt, end_dt))
        if not dates:
            raise ValueError("No dates found.")
        # Limit batch size to prevent abuse
        if len(dates) > 52:
            raise ValueError("Batch too large. Maximum 52 weeks per batch.")
    except Exception as e:
        emit("error", {"msg": str(e)})
        return

    def task():
        used_themes = []
        used_descs = []
        acquired = _task_semaphore.acquire(timeout=5)
        if not acquired:
            emit_log(sid, "[Batch] Server busy. Please try again later.")
            socketio.emit("task_done", {"task": "batch"}, to=sid)
            return
        try:
            from playwright.sync_api import sync_playwright
            emit_log(sid, f"[Batch] Starting: {len(dates)} weekly records.")

            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True, slow_mo=60)
                page = browser.new_page()
                page.set_default_timeout(PLAYWRIGHT_TIMEOUT_MS)
                page.set_default_navigation_timeout(PLAYWRIGHT_TIMEOUT_MS)
                login_and_wait_home(page, user, pw)
                record_list_ctx = open_records_list_ctx(page)

                total = len(dates)
                for idx, dt_item in enumerate(dates, start=1):
                    date_ymd = f"{dt_item.year:04d}/{dt_item.month:02d}/{dt_item.day:02d}"
                    emit_log(sid, f"[Batch] ({idx}/{total}) Generating theme + desc for {date_ymd}...")

                    theme, desc = generate_weekly_theme_desc_deepseek(
                        api_key=LLM_API_KEY, club_name=club, date_ymd=date_ymd,
                        club_desc=club_desc, periodic_desc=periodic,
                        used_themes=used_themes, used_descs=used_descs,
                    )
                    if not theme or not desc:
                        raise RuntimeError(f"Qwen returned empty for {date_ymd}.")
                    used_themes.append(theme)
                    used_descs.append(desc)
                    socketio.emit("preview_record", {"text": f"{theme}\n\n{desc}"}, to=sid)
                    emit_log(sid, f"[Batch] ({idx}/{total}) Filling record...")

                    add_ctx = open_add_record_ctx(record_list_ctx, page)
                    select_club_by_text(add_ctx, club)

                    date_input = add_ctx.locator("div.layui-form-item:has(label:has-text('Event date')) input")
                    date_input.click()
                    cal_scope = add_ctx if add_ctx.locator("#layui-laydate1").count() else page
                    select_date_layui(cal_scope, dt_item.year, dt_item.month, dt_item.day)

                    add_ctx.locator("div.layui-form-item:has(label:has-text('Activity theme')) input").fill(theme)
                    add_ctx.locator("input[name='CDuration']").fill(c)
                    add_ctx.locator("input[name='ADuration']").fill(a)
                    add_ctx.locator("input[name='SDuration']").fill(s)
                    add_ctx.locator("textarea[name='Reflection']").fill(desc)
                    add_ctx.locator("button[lay-filter='add']:has-text('Save')").click()
                    emit_log(sid, f"[Batch] ({idx}/{total}) Save clicked.")

                    try:
                        page.locator("iframe[src*='/Stu/Cas/AddRecord']").wait_for(state="detached")
                    except Exception:
                        time.sleep(1.2)

                browser.close()

            emit_log(sid, "[Batch] Run finished.")
            socketio.emit("task_done", {"task": "batch"}, to=sid)
        except Exception as e:
            emit_log(sid, f"[Batch] Error: {sanitize_error(e)}")
            socketio.emit("task_done", {"task": "batch"}, to=sid)
        finally:
            _task_semaphore.release()

    threading.Thread(target=task, daemon=True).start()


@socketio.on("run_reflection")
def handle_run_reflection(data):
    sid = request.sid
    user = safe_str(data.get("username", ""), max_length=100)
    pw = safe_str(data.get("password", ""), max_length=200)
    club = safe_str(data.get("club", ""), max_length=200)
    club_desc = safe_str(data.get("club_desc", ""), max_length=1000)
    raw_titles = data.get("titles", [])
    raw_desc_lines = data.get("desc_lines", [])
    raw_selected = data.get("outcomes", [])
    # Optional pre-generated (user-approved) content from the chat agent.
    raw_summaries = data.get("summaries", [])
    raw_contents = data.get("contents", [])

    if not isinstance(raw_titles, list):
        emit("error", {"msg": "Invalid titles format."})
        return
    titles = [safe_str(t, max_length=500) for t in raw_titles if isinstance(t, str) and t.strip()]
    desc_lines = [safe_str(d, max_length=1000) for d in raw_desc_lines] if isinstance(raw_desc_lines, list) else []
    selected = [s for s in raw_selected if isinstance(s, str) and s in VALID_OUTCOMES]
    approved_summaries = [safe_str(x, max_length=2000) for x in raw_summaries] if isinstance(raw_summaries, list) else []
    approved_contents = [safe_str(x, max_length=20000) for x in raw_contents] if isinstance(raw_contents, list) else []
    # Only trust approved content if it covers every title.
    use_approved = len(approved_summaries) == len(titles) and len(approved_contents) == len(titles) and len(titles) > 0

    if not titles:
        emit("error", {"msg": "Please provide at least one title."})
        return
    if len(titles) > 20:
        emit("error", {"msg": "Maximum 20 reflections per batch."})
        return
    if not selected:
        emit("error", {"msg": "Select at least one Learning Outcome."})
        return
    if not desc_lines:
        desc_lines = [""] * len(titles)

    def task():
        acquired = _task_semaphore.acquire(timeout=5)
        if not acquired:
            emit_log(sid, "[Reflection] Server busy. Please try again later.")
            socketio.emit("task_done", {"task": "reflection"}, to=sid)
            return
        try:
            from playwright.sync_api import sync_playwright
            emit_log(sid, f"[Reflection] Starting: {len(titles)} reflections.")

            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True, slow_mo=60)
                page = browser.new_page()
                page.set_default_timeout(PLAYWRIGHT_TIMEOUT_MS)
                page.set_default_navigation_timeout(PLAYWRIGHT_TIMEOUT_MS)
                login_and_wait_home(page, user, pw)
                refl_list_ctx = open_reflection_list_ctx(page)

                total = len(titles)
                for idx, title in enumerate(titles, start=1):
                    reflection_desc = desc_lines[idx - 1] if idx - 1 < len(desc_lines) else ""
                    emit_log(sid, f"[Reflection] ({idx}/{total}) Opening add dialog...")
                    add_ctx = open_add_reflection_ctx(refl_list_ctx, page)

                    emit_log(sid, f"[Reflection] ({idx}/{total}) Selecting club: {club}")
                    select_club_by_text(add_ctx, club)
                    add_ctx.locator("input[name='Title']").fill(title)

                    if use_approved:
                        summary = approved_summaries[idx - 1]
                        reflection_text = approved_contents[idx - 1]
                        emit_log(sid, f"[Reflection] ({idx}/{total}) Using approved content.")
                    else:
                        emit_log(sid, f"[Reflection] ({idx}/{total}) Generating summary...")
                        summary = generate_reflection_summary_deepseek(
                            api_key=LLM_API_KEY, club_name=club, title=title,
                            club_desc=club_desc, reflection_desc=reflection_desc,
                        )
                        emit_log(sid, f"[Reflection] ({idx}/{total}) Generating reflection content...")
                        reflection_text = generate_reflection_content_deepseek(
                            api_key=LLM_API_KEY, club_name=club, title=title,
                            club_desc=club_desc, reflection_desc=reflection_desc,
                        )
                    socketio.emit("preview_reflection", {"summary": summary, "content": reflection_text}, to=sid)

                    add_ctx.locator("textarea[name='Summary']").fill(summary)
                    fill_kindeditor_body(add_ctx, reflection_text)

                    emit_log(sid, f"[Reflection] ({idx}/{total}) Selecting outcomes: {', '.join(selected)}")
                    click_learning_outcomes(add_ctx, selected)

                    add_ctx.locator("button[lay-filter='add']:has-text('Save')").click()
                    emit_log(sid, f"[Reflection] ({idx}/{total}) Save clicked.")

                    try:
                        page.locator("iframe[src*='/Stu/Cas/AddReflection']").wait_for(state="detached")
                    except Exception:
                        time.sleep(1.2)

                browser.close()

            emit_log(sid, "[Reflection] Run finished.")
            socketio.emit("task_done", {"task": "reflection"}, to=sid)
        except Exception as e:
            emit_log(sid, f"[Reflection] Error: {sanitize_error(e)}")
            socketio.emit("task_done", {"task": "reflection"}, to=sid)
        finally:
            _task_semaphore.release()

    threading.Thread(target=task, daemon=True).start()


# --------------- CAS Agent (function calling) ---------------

VALID_OUTCOMES = {"Awareness", "Challenge", "Initiative", "Collaboration",
                  "Commitment", "Global Value", "Ethics", "New Skills"}

WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# OpenAI-compatible tool schemas the model can call to drive autofill.
CAS_TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "create_activity_record",
            "description": "Create ONE IB CAS activity record for a single date. Use when the user wants to log a single activity on a specific day.",
            "parameters": {
                "type": "object",
                "properties": {
                    "club": {"type": "string", "description": "Club name. Must match one of the available clubs exactly."},
                    "date": {"type": "string", "description": "Activity date in YYYY/MM/DD format. Required; never default to today's date if the user did not explicitly provide a date."},
                    "theme": {"type": "string", "description": "Short activity theme, e.g. 'Cold War Origins Discussion'."},
                    "c_hours": {"type": "string", "description": "Creativity hours, a number. Use '0' only for categories the user omitted after they explicitly provided at least one C/A/S hour."},
                    "a_hours": {"type": "string", "description": "Activity hours, a number. Use '0' only for categories the user omitted after they explicitly provided at least one C/A/S hour."},
                    "s_hours": {"type": "string", "description": "Service hours, a number. Use '0' only for categories the user omitted after they explicitly provided at least one C/A/S hour."},
                },
                "required": ["club", "date", "theme"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_weekly_batch",
            "description": "Create MULTIPLE weekly IB CAS activity records across a date range (same weekday each week). Use when the user wants recurring weekly records over weeks/months.",
            "parameters": {
                "type": "object",
                "properties": {
                    "club": {"type": "string", "description": "Club name. Must match one of the available clubs exactly."},
                    "club_desc": {"type": "string", "description": "Brief description of the weekly activity or club, based on details the user explicitly provided."},
                    "weekday": {"type": "string", "enum": WEEKDAYS, "description": "The weekday on which the activity recurs."},
                    "start_date": {"type": "string", "description": "First date in YYYY/MM/DD. Must fall on the given weekday."},
                    "end_date": {"type": "string", "description": "Last date in YYYY/MM/DD. Must fall on the given weekday."},
                    "periodic": {"type": "string", "description": "Optional. A recurring overarching activity all weeks share."},
                    "c_hours": {"type": "string", "description": "Creativity hours per record. Use '0' only for categories the user omitted after they explicitly provided at least one C/A/S hour."},
                    "a_hours": {"type": "string", "description": "Activity hours per record. Use '0' only for categories the user omitted after they explicitly provided at least one C/A/S hour."},
                    "s_hours": {"type": "string", "description": "Service hours per record. Use '0' only for categories the user omitted after they explicitly provided at least one C/A/S hour."},
                },
                "required": ["club", "club_desc", "weekday", "start_date", "end_date"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_reflection",
            "description": "Create one or more IB CAS activity reflections. Use when the user wants to write reflections with learning outcomes.",
            "parameters": {
                "type": "object",
                "properties": {
                    "club": {"type": "string", "description": "Club name. Must match one of the available clubs exactly."},
                    "club_desc": {"type": "string", "description": "Brief description of the club."},
                    "titles": {"type": "array", "items": {"type": "string"}, "description": "One title per reflection."},
                    "desc_lines": {"type": "array", "items": {"type": "string"}, "description": "Optional. One focus description per reflection, aligned with titles."},
                    "outcomes": {"type": "array", "items": {"type": "string", "enum": sorted(VALID_OUTCOMES)}, "description": "Learning outcomes to tick for every reflection."},
                },
                "required": ["club", "titles", "outcomes"],
            },
        },
    },
]


def _norm_hours(v, default="0"):
    v = str(v).strip() if v is not None else ""
    return v if re.match(r"^\d+(\.\d+)?$", v) else (default if not v else v)


DATE_HINT_RE = re.compile(
    r"\b\d{4}[/-]\d{1,2}[/-]\d{1,2}\b|"
    r"\b\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?\b|"
    r"\b(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
    r"jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|"
    r"dec(?:ember)?)\s+\d{1,2}(?:,\s*\d{4})?\b|"
    r"\b(?:today|tomorrow|yesterday)\b|"
    r"\d{1,2}\s*\u6708\s*\d{1,2}\s*\u65e5|"
    r"\u4eca\u5929|\u660e\u5929|\u6628\u5929",
    re.I,
)
HOURS_HINT_RE = re.compile(
    r"\b[CAS]\s*[:=]?\s*\d+(?:\.\d+)?\s*(?:h|hr|hrs|hour|hours)?\b|"
    r"\b(?:creativity|activity|service)\s*(?:hours?)?\s*[:=]?\s*\d+(?:\.\d+)?\b|"
    r"\b\d+(?:\.\d+)?\s*(?:h|hr|hrs|hour|hours)\s*(?:of\s*)?(?:creativity|activity|service)\b|"
    r"[CAS]\s*\d+(?:\.\d+)?\s*\u5c0f\u65f6|"
    r"(?:\u521b\u9020|\u521b\u9020\u529b|\u6d3b\u52a8|\u670d\u52a1)\s*\d+(?:\.\d+)?\s*\u5c0f\u65f6",
    re.I,
)
_GENERIC_CONTENT_TOKENS = {
    "activity", "record", "reflection", "club", "cas", "session", "meeting",
    "event", "experience", "write", "create", "fill", "log", "make", "about",
    "weekly", "batch", "semester", "school", "student", "students",
}


def _looks_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def _agent_context_text(history, user_message: str, attachments=None) -> str:
    parts = [user_message or ""]
    if isinstance(history, list):
        for h in history:
            if isinstance(h, dict) and h.get("role") == "user" and isinstance(h.get("content"), str):
                parts.append(h["content"])
    for nm, tx in attachments or []:
        parts.append(nm or "")
        parts.append(tx or "")
    return "\n".join(parts)


def _has_explicit_hours(context: str) -> bool:
    return bool(HOURS_HINT_RE.search(context or ""))


def _context_tokens(text: str) -> list:
    return re.findall(r"[a-z0-9]+|[\u4e00-\u9fff]+", (text or "").lower())


def _meaningful_tokens(text: str, exclude: str = "") -> set:
    exclude_tokens = set(_context_tokens(exclude))
    tokens = set()
    for tok in _context_tokens(text):
        if tok in exclude_tokens or tok in _GENERIC_CONTENT_TOKENS:
            continue
        if re.fullmatch(r"\d+(?:\.\d+)?", tok):
            continue
        if re.fullmatch(r"[a-z0-9]+", tok) and len(tok) <= 2:
            continue
        tokens.add(tok)
    return tokens


def _value_supported_by_context(value: str, context: str, exclude: str = "") -> bool:
    value = safe_str(value, 1000)
    if not value:
        return False
    value_norm = re.sub(r"\s+", " ", value).strip().lower()
    context_norm = re.sub(r"\s+", " ", context or "").strip().lower()
    if value_norm and value_norm in context_norm:
        return True
    value_tokens = _meaningful_tokens(value, exclude)
    if not value_tokens:
        return False
    context_tokens = _meaningful_tokens(context, exclude)
    overlap = value_tokens & context_tokens
    required = 1 if len(value_tokens) <= 2 else max(2, (len(value_tokens) + 1) // 2)
    return len(overlap) >= required


def _has_supported_content(values, context: str, exclude: str = "") -> bool:
    for value in values:
        if _value_supported_by_context(value, context, exclude):
            return True
    return False


def _club_is_valid(club: str, clubs: list) -> bool:
    return bool(club) and (not clubs or club in clubs)


def _missing_followup(action: str, missing: list, user_message: str) -> str:
    zh = _looks_chinese(user_message)
    labels = {
        "club": ("\u793e\u56e2", "club"),
        "valid_club": ("\u53ef\u7528\u793e\u56e2\u5217\u8868\u4e2d\u7684\u793e\u56e2", "a club from the available club list"),
        "loaded_clubs": ("\u8bf7\u5148\u767b\u5f55\u5e76\u5237\u65b0\u793e\u56e2", "please sign in and fetch clubs first"),
        "date": ("\u6d3b\u52a8\u65e5\u671f", "activity date"),
        "date_range": ("\u5f00\u59cb\u548c\u7ed3\u675f\u65e5\u671f", "start and end dates"),
        "activity_detail": ("\u6d3b\u52a8\u4e3b\u9898\u6216\u6d3b\u52a8\u63cf\u8ff0", "activity theme or description"),
        "reflection_detail": ("\u53cd\u601d\u5185\u5bb9\u6216\u7ecf\u5386\u80cc\u666f", "reflection focus or experience details"),
        "batch_detail": ("\u6bcf\u5468\u6d3b\u52a8\u5185\u5bb9\u6216\u793e\u56e2\u63cf\u8ff0", "weekly activity details or club description"),
        "hours": ("C/A/S \u5c0f\u65f6", "C/A/S hours"),
    }
    items = [labels.get(m, (m, m))[0 if zh else 1] for m in missing]
    if missing == ["loaded_clubs"]:
        return items[0] + "\u3002" if zh else items[0].capitalize() + "."
    if len(items) == 1:
        return ("\u8bf7\u544a\u8bc9\u6211" + items[0] + "\u3002") if zh else f"Please tell me the {items[0]}."
    return ("\u8bf7\u8865\u5145\uff1a" + "\u3001".join(items) + "\u3002") if zh else "Please tell me: " + ", ".join(items) + "."


def _missing_required_details(name: str, args: dict, history, user_message: str, attachments, clubs: list) -> list:
    context = _agent_context_text(history, user_message, attachments)
    club = safe_str(args.get("club", ""), 200)
    if not clubs:
        return ["loaded_clubs"]

    missing = []
    if not club:
        missing.append("club")
    elif not _club_is_valid(club, clubs):
        missing.append("valid_club")

    if name == "create_activity_record":
        if not DATE_HINT_RE.search(context):
            missing.append("date")
        content_values = [args.get("activity_desc", ""), args.get("theme", "")]
        if not _has_supported_content(content_values, context, club):
            missing.append("activity_detail")
        if not _has_explicit_hours(context):
            missing.append("hours")
    elif name == "create_reflection":
        content_values = []
        for key in ("club_desc",):
            content_values.append(args.get(key, ""))
        for key in ("titles", "desc_lines"):
            raw = args.get(key, [])
            if isinstance(raw, list):
                content_values.extend(x for x in raw if isinstance(x, str))
        if not _has_supported_content(content_values, context, club):
            missing.append("reflection_detail")
    elif name == "create_weekly_batch":
        if len(DATE_HINT_RE.findall(context)) < 2:
            missing.append("date_range")
        content_values = [args.get("club_desc", ""), args.get("periodic", "")]
        if not _has_supported_content(content_values, context, club):
            missing.append("batch_detail")
        if not _has_explicit_hours(context):
            missing.append("hours")
    return missing


def _chat_history_messages_for_budget(history, char_budget: int) -> list:
    if not isinstance(history, list) or char_budget <= 0:
        return []
    picked = []
    used = 0
    for h in reversed(history):
        if not isinstance(h, dict) or h.get("role") not in ("user", "assistant"):
            continue
        if not isinstance(h.get("content"), str):
            continue
        content = safe_str(h.get("content", ""), max_length=CHAT_MAX_MESSAGE_CHARS).strip()
        if not content:
            continue
        cost = len(content)
        if used + cost > char_budget:
            remaining = char_budget - used
            if remaining > 200:
                picked.append({"role": h["role"], "content": content[-remaining:]})
            break
        picked.append({"role": h["role"], "content": content})
        used += cost
    picked.reverse()
    return picked


def build_record_proposal(args: dict) -> dict:
    """Validate args, pre-generate the description, and return an approvable proposal."""
    thinking_enabled = bool(args.get("thinking_enabled", False))
    club = safe_str(args.get("club", ""), 200)
    date_str = safe_str(args.get("date", ""), 20).replace("-", "/")
    theme = safe_str(args.get("theme", ""), 500)
    activity_desc = safe_str(args.get("activity_desc", ""), 1000)
    c = _norm_hours(args.get("c_hours", "0"))
    a = _norm_hours(args.get("a_hours", "0"))
    s = _norm_hours(args.get("s_hours", "0"))
    y, mo, d = parse_date_ymd(date_str)
    date_ymd = f"{y:04d}/{mo:02d}/{d:02d}"
    for nm, val in [("C", c), ("A", a), ("S", s)]:
        if not re.match(r"^\d+(\.\d+)?$", val):
            raise ValueError(f"{nm} hours must be a number.")
    if not theme and activity_desc:
        theme = generate_activity_title_deepseek(
            api_key=LLM_API_KEY, club_name=club, date_ymd=date_ymd,
            activity_desc=activity_desc, thinking_enabled=thinking_enabled,
        )
    if not theme:
        raise ValueError("Activity theme is required.")
    description = generate_activity_record_from_context_deepseek(
        api_key=LLM_API_KEY, club_name=club, date_ymd=date_ymd,
        theme=theme, c_hours=c, a_hours=a, s_hours=s,
        activity_desc=activity_desc, thinking_enabled=thinking_enabled,
    )
    return {
        "action": "record",
        "params": {"club": club, "club_display": english_club_name(club), "date": date_ymd,
                   "theme": theme, "activity_desc": activity_desc,
                   "c_hours": c, "a_hours": a, "s_hours": s},
        "content": {"description": description},
    }


def build_reflection_proposal(args: dict) -> dict:
    thinking_enabled = bool(args.get("thinking_enabled", False))
    club = safe_str(args.get("club", ""), 200)
    club_desc = safe_str(args.get("club_desc", ""), 1000)
    raw_titles = args.get("titles", [])
    raw_descs = args.get("desc_lines", [])
    raw_outcomes = args.get("outcomes", [])
    titles = [safe_str(t, 500) for t in raw_titles if isinstance(t, str) and t.strip()]
    desc_lines = [safe_str(x, 1000) for x in raw_descs] if isinstance(raw_descs, list) else []
    outcomes = [o for o in raw_outcomes if isinstance(o, str) and o in VALID_OUTCOMES]
    if not titles:
        raise ValueError("At least one reflection title is required.")
    if len(titles) > 10:
        raise ValueError("Maximum 10 reflections per chat proposal.")
    if not outcomes:
        raise ValueError("At least one Learning Outcome is required.")
    if not desc_lines:
        desc_lines = [""] * len(titles)
    summaries, contents = [], []
    for i, title in enumerate(titles):
        rdesc = desc_lines[i] if i < len(desc_lines) else ""
        summaries.append(generate_reflection_summary_deepseek(
            api_key=LLM_API_KEY, club_name=club, title=title,
            club_desc=club_desc, reflection_desc=rdesc,
            thinking_enabled=thinking_enabled))
        contents.append(generate_reflection_content_deepseek(
            api_key=LLM_API_KEY, club_name=club, title=title,
            club_desc=club_desc, reflection_desc=rdesc,
            thinking_enabled=thinking_enabled))
    return {
        "action": "reflection",
        "params": {"club": club,
                   "club_display": ("Overall CAS Reflection (Conversation)"
                                    if club.strip() == CONVERSATION_CLUB else english_club_name(club)),
                   "club_desc": club_desc,
                   "titles": titles, "desc_lines": desc_lines, "outcomes": outcomes},
        "content": {"summaries": summaries, "contents": contents},
    }


def build_batch_proposal(args: dict) -> dict:
    club = safe_str(args.get("club", ""), 200)
    club_desc = safe_str(args.get("club_desc", ""), 1000)
    weekday = safe_str(args.get("weekday", ""), 20)
    start = safe_str(args.get("start_date", ""), 20).replace("-", "/")
    end = safe_str(args.get("end_date", ""), 20).replace("-", "/")
    periodic = safe_str(args.get("periodic", ""), 500)
    c = _norm_hours(args.get("c_hours", "0"))
    a = _norm_hours(args.get("a_hours", "0"))
    s = _norm_hours(args.get("s_hours", "0"))
    if weekday not in WEEKDAYS:
        raise ValueError("Invalid weekday.")
    y1, m1, d1 = parse_date_ymd(start)
    y2, m2, d2 = parse_date_ymd(end)
    start_dt = dt_date(y1, m1, d1)
    end_dt = dt_date(y2, m2, d2)
    widx = WEEKDAYS.index(weekday)
    if start_dt.weekday() != widx or end_dt.weekday() != widx:
        raise ValueError(f"Start and end dates must both fall on {weekday}.")
    if start_dt > end_dt:
        raise ValueError("End date must be after start date.")
    dates = list(iter_weekly_dates(start_dt, end_dt))
    if not dates:
        raise ValueError("No dates found in range.")
    if len(dates) > 52:
        raise ValueError("Batch too large. Maximum 52 weeks.")
    # Generate a sample (first week) so the user sees the style before approving.
    sample_theme, sample_desc = generate_weekly_theme_desc_deepseek(
        api_key=LLM_API_KEY, club_name=club,
        date_ymd=f"{dates[0].year:04d}/{dates[0].month:02d}/{dates[0].day:02d}",
        club_desc=club_desc, periodic_desc=periodic, used_themes=[], used_descs=[])
    return {
        "action": "batch",
        "params": {"club": club, "club_display": english_club_name(club), "club_desc": club_desc, "weekday": weekday,
                   "start_date": f"{y1:04d}/{m1:02d}/{d1:02d}",
                   "end_date": f"{y2:04d}/{m2:02d}/{d2:02d}",
                   "periodic": periodic, "c_hours": c, "a_hours": a, "s_hours": s},
        "content": {"week_count": len(dates),
                    "sample_theme": sample_theme, "sample_desc": sample_desc},
    }


_PROPOSAL_BUILDERS = {
    "create_activity_record": build_record_proposal,
    "create_weekly_batch": build_batch_proposal,
    "create_reflection": build_reflection_proposal,
}

# Detects when the assistant *says* it will act but forgot to actually call a tool
# (e.g. "let me create the record now", "I'll proceed", "让我现在创建记录").
_INTENT_EN = re.compile(
    r"(let me(?! know)|i['’]?ll|i will|i'?m going to|i am going to|i'?ll go ahead|let me go ahead|"
    r"i'?ll proceed|let me proceed)\b"
    r"[^.!?\n]{0,80}?\b(creat|fill|add|submit|log|make|proceed|generat|draft|set up|set a)",
    re.I,
)
_INTENT_ZH = re.compile(
    r"(让我|我来|我帮你?|我现在就|这就)[^。！？\n]{0,40}?(创建|填写|提交|添加|生成|帮你填|做一个|起草|记录|反思)"
)


def _first_tool_call(msg: dict):
    tcs = msg.get("tool_calls") or []
    return tcs[0] if tcs else None


def _force_tool_call(base_messages: list, prior_text: str, thinking_enabled: Optional[bool] = None):
    """Re-prompt the model to actually emit a function call when it stalled.

    Tries tool_choice='required' (forces a call) and falls back to 'auto' with an
    explicit instruction. Returns a tool_call dict or None. Best-effort: any API
    error is swallowed so the caller can fall back to plain text.
    """
    nudged = base_messages + [
        {"role": "assistant", "content": prior_text or "Okay."},
        {"role": "user", "content": (
            "If all required details are explicitly present, do NOT reply with text; "
            "call the correct function now. If a required detail is missing, ask for "
            "that exact detail instead. Never invent or default dates.")},
    ]
    choices = ("required", "auto")
    for choice in choices:
        try:
            r = llm_chat(LLM_API_KEY, LLM_MODEL, nudged,
                              temperature=0.1, max_tokens=600,
                              tools=CAS_TOOLS, tool_choice=choice,
                              thinking_enabled=False)
            call = _first_tool_call(r["choices"][0]["message"])
            if call:
                return call
        except Exception:
            continue
    return None


# --------------- File upload (Qwen-Long file-extract) ---------------

def extract_file_via_qwen(filename: str, data: bytes) -> dict:
    """Upload a file to Qwen-Long and return extracted text.

    Qwen-Long supports file-extract uploads and fileid:// references. Best-effort:
    the uploaded file is deleted afterwards so storage does not accumulate.
    Returns {name?, text, note, chars, truncated}.
    """
    fail = lambda note: {"text": "", "note": note, "chars": 0, "truncated": False}
    if not LLM_API_KEY:
        return fail("Server API key not configured.")
    headers = {"Authorization": f"Bearer {LLM_API_KEY}"}
    file_id = None
    try:
        resp = requests.post(
            LLM_FILES_ENDPOINT, headers=headers,
            files={"file": (filename, io.BytesIO(data))},
            data={"purpose": "file-extract"}, timeout=120,
        )
        if resp.status_code != 200:
            return fail(f"Upload failed (HTTP {resp.status_code}).")
        file_id = (resp.json() or {}).get("id")
        if not file_id:
            return fail("Upload returned no file id.")

        messages = [
            {"role": "system", "content": (
                "You extract readable text from uploaded files for a CAS writing assistant. "
                "Output only the extracted useful text. Do not summarize unless the source is "
                "an image/photo with no exact text; in that case describe the visible content "
                "factually and concisely.")},
            {"role": "system", "content": f"fileid://{file_id}"},
            {"role": "user", "content": (
                f"Extract the useful text/content from {filename}. Preserve names, dates, "
                "activity details, and table rows that may help write an IB CAS record or reflection.")}
        ]
        text = ""
        last_note = ""
        for _ in range(6):
            try:
                rc = llm_chat(LLM_API_KEY, FILE_MODEL, messages,
                              temperature=0.1, max_tokens=4096,
                              thinking_enabled=False)
                text = llm_message_content(rc)
                break
            except Exception as e:
                last_note = sanitize_error(e)
                if "File parsing in progress" not in last_note and "parsing" not in last_note.lower():
                    return fail(f"Extraction failed: {last_note}")
                time.sleep(2)
        if not text and last_note:
            return fail(f"Extraction failed: {last_note}")
    except Exception as e:
        return fail(f"Could not read file: {sanitize_error(e)}")
    finally:
        if file_id:
            try:
                requests.delete(f"{LLM_FILES_ENDPOINT}/{file_id}", headers=headers, timeout=30)
            except Exception:
                pass

    text = (text or "").strip()
    truncated = len(text) > MAX_EXTRACT_CHARS
    if truncated:
        text = text[:MAX_EXTRACT_CHARS]
    note = "" if text else "No text could be extracted (an image may contain no readable text)."
    return {"text": text, "note": note, "chars": len(text), "truncated": truncated}


# --------------- Routes ---------------

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/fonts/<path:filename>")
def font_file(filename):
    fonts_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
    return send_from_directory(fonts_dir, filename)


@app.route("/api/upload", methods=["POST"])
def upload_files():
    client_ip = request.remote_addr or "unknown"
    if not _check_rate_limit(client_ip):
        return jsonify({"error": "Rate limit exceeded. Please wait before uploading more."}), 429
    if not LLM_API_KEY:
        return jsonify({"error": "Qwen/DashScope API key not configured on server."}), 500

    files = request.files.getlist("files")
    if not files:
        return jsonify({"error": "No files uploaded."}), 400
    if len(files) > MAX_UPLOAD_FILES:
        return jsonify({"error": f"Too many files (max {MAX_UPLOAD_FILES})."}), 400

    results = []
    for f in files:
        filename = safe_str(f.filename or "file", 255) or "file"
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in ALLOWED_UPLOAD_EXT:
            results.append({"name": filename, "text": "", "note": "Unsupported file type.",
                            "chars": 0, "truncated": False})
            continue
        data = f.read(MAX_UPLOAD_BYTES + 1)
        if len(data) > MAX_UPLOAD_BYTES:
            results.append({"name": filename, "text": "",
                            "note": f"File too large (max {MAX_UPLOAD_BYTES // (1024 * 1024)} MB).",
                            "chars": 0, "truncated": False})
            continue
        info = extract_file_via_qwen(filename, data)
        info["name"] = filename
        results.append(info)
    return jsonify({"files": results})


def quick_attachment_context(raw_attachments) -> str:
    if not isinstance(raw_attachments, list):
        return ""
    blocks = []
    for a in raw_attachments[:MAX_UPLOAD_FILES]:
        if not isinstance(a, dict):
            continue
        name = safe_str(a.get("name", ""), 255) or "file"
        text = safe_str(a.get("text", ""), 3000)
        if text:
            blocks.append(f"[Attached file/photo: {name}]\n{text}")
    return "\n\n".join(blocks)


@app.route("/api/quick_proposals", methods=["POST"])
def quick_proposals():
    client_ip = request.remote_addr or "unknown"
    if not _check_rate_limit(client_ip):
        return jsonify({"error": "Rate limit exceeded. Please wait before generating more forms."}), 429

    data = request.json or {}
    kind = safe_str(data.get("kind", ""), 30)
    thinking_enabled = bool(data.get("thinking", False))
    raw_requests = data.get("requests", [])
    if kind not in ("record", "reflection"):
        return jsonify({"error": "Invalid request type."}), 400
    if not isinstance(raw_requests, list) or not raw_requests:
        return jsonify({"error": "No requests provided."}), 400
    if len(raw_requests) > 10:
        return jsonify({"error": "Maximum 10 requests at a time."}), 400
    if not LLM_API_KEY:
        return jsonify({"error": "Qwen/DashScope API key not configured on server."}), 500

    proposals = []
    try:
        for raw in raw_requests:
            if not isinstance(raw, dict):
                raise ValueError("Invalid request item.")
            club = safe_str(raw.get("club", ""), 200)
            desc = safe_str(raw.get("description", ""), 2000)
            attach_context = quick_attachment_context(raw.get("attachments", []))
            full_desc = "\n\n".join(x for x in [desc, attach_context] if x).strip()
            if not club:
                raise ValueError("Please select a club for every request.")
            if not full_desc:
                raise ValueError("Please enter a description or attach at least one readable file/photo for every request.")

            if kind == "record":
                date_str = safe_str(raw.get("date", ""), 20).replace("-", "/")
                if not date_str:
                    raise ValueError("Please select a date for every record.")
                y, mo, d = parse_date_ymd(date_str)
                date_ymd = f"{y:04d}/{mo:02d}/{d:02d}"
                theme = generate_activity_title_deepseek(
                    api_key=LLM_API_KEY, club_name=club, date_ymd=date_ymd,
                    activity_desc=full_desc, thinking_enabled=thinking_enabled,
                )
                proposals.append(build_record_proposal({
                    "club": club,
                    "date": date_ymd,
                    "theme": theme,
                    "activity_desc": full_desc,
                    "c_hours": raw.get("c_hours", "0"),
                    "a_hours": raw.get("a_hours", "0"),
                    "s_hours": raw.get("s_hours", "0"),
                    "thinking_enabled": thinking_enabled,
                }))
            else:
                outcomes = raw.get("outcomes", [])
                if not isinstance(outcomes, list):
                    outcomes = []
                outcomes = [o for o in outcomes if isinstance(o, str) and o in VALID_OUTCOMES]
                if not outcomes:
                    raise ValueError("Please select at least one Learning Outcome for every reflection.")
                title = generate_reflection_title_deepseek(
                    api_key=LLM_API_KEY, club_name=club, reflection_desc=full_desc,
                    thinking_enabled=thinking_enabled,
                )
                proposals.append(build_reflection_proposal({
                    "club": club,
                    "club_desc": safe_str(raw.get("club_desc", ""), 1000),
                    "titles": [title],
                    "desc_lines": [full_desc],
                    "outcomes": outcomes,
                    "thinking_enabled": thinking_enabled,
                }))
    except ValueError as ve:
        return jsonify({"error": str(ve)}), 400
    except Exception as e:
        return jsonify({"error": sanitize_error(e)}), 500

    return jsonify({"proposals": proposals})


@app.route("/api/chat", methods=["POST"])
def chat_cas():
    # Rate limiting
    client_ip = request.remote_addr or "unknown"
    if not _check_rate_limit(client_ip):
        return jsonify({"error": "Rate limit exceeded. Please wait before sending more messages."}), 429

    data = request.json or {}
    user_message = safe_str(data.get("message", ""), max_length=CHAT_MAX_MESSAGE_CHARS)
    history = data.get("history", [])
    thinking_enabled = bool(data.get("thinking", False))
    raw_clubs = data.get("clubs", [])
    clubs = [safe_str(c, 200) for c in raw_clubs if isinstance(c, str) and c.strip()][:200] \
        if isinstance(raw_clubs, list) else []

    # Attachments: text already extracted client-side via /api/upload (Qwen-Long file-extract).
    raw_attach = data.get("attachments", [])
    attachments = []
    if isinstance(raw_attach, list):
        for a in raw_attach[:MAX_UPLOAD_FILES]:
            if isinstance(a, dict):
                nm = safe_str(a.get("name", ""), 255)
                tx = safe_str(a.get("text", ""), MAX_EXTRACT_CHARS)
                if tx:
                    attachments.append((nm or "file", tx))

    if not user_message and not attachments:
        return jsonify({"error": "Empty message"}), 400
    if not user_message:
        user_message = "Please read the attached file(s) and respond."
    if not LLM_API_KEY:
        return jsonify({"error": "Qwen/DashScope API key not configured on server."}), 500

    # Validate history is a list of dicts
    if not isinstance(history, list):
        history = []

    today = dt_date.today().strftime("%Y/%m/%d")
    system_content = (
        "You are an IB CAS (Creativity, Activity, Service) expert advisor AND autofill agent for "
        "WFLA high school students. You can answer CAS questions, and you can also fill in the school "
        "system on the student's behalf by calling functions.\n\n"
        f"Today's date is {today}.\n\n"
        "Behaviour rules:\n"
        "- For general questions, just answer in plain text. Be concise and supportive.\n"
        "- When the user asks to log/create/fill a record, reflection, or weekly batch and you have "
        "the required details, you MUST call the matching function IN THIS SAME TURN.\n"
        "- CRITICAL: Never announce that you will create something without calling the function. Do "
        "NOT reply with phrases like 'let me create the record', 'I'll create it now', or 'creating "
        "now'. If you are about to write such a sentence, call the function instead — the app shows "
        "the user an approval card automatically, so you never need to ask for confirmation in text.\n"
        "- Ask a short follow-up question when required details are missing or ambiguous. Do not "
        "guess, invent, or silently fill required details.\n"
        "- Activity records require: a valid club, activity date, activity theme/description, and "
        "at least one explicit C/A/S hour value. Use 0 only for C/A/S categories the user omitted "
        "after giving at least one hour value. NEVER default a missing record date to today's date.\n"
        "- Reflections require: a valid club (or Conversation when available) and enough experience "
        "details to write a real reflection. You may infer outcomes and create a concise title, but "
        "do not invent the student's experience.\n"
        "- Weekly batch records require: a valid club, start and end dates, weekly activity details "
        "or club description, and at least one explicit C/A/S hour value.\n"
        "- You may briefly note a concern (e.g. a CAS rule), but you must STILL call the function in "
        "the same turn unless a required detail is missing.\n"
        "- Convert any natural-language date to YYYY/MM/DD. Resolve relative dates using today's date.\n"
        "- Reply in the same language the user uses (Chinese or English).\n"
        "- Never invent a club name. Only use clubs from the available list below.\n"
        "- If the user gives an English club alias such as 'Computerization Club', match it to the "
        "available club whose English name appears in parentheses, e.g. '世外信息化社(Computerization)'. "
        "Use the exact available club name in the function call.\n"
        "- For reflections, if the user did not specify Learning Outcomes, infer 3-5 suitable "
        "outcomes from the activity description. Do NOT ask only because outcomes are missing.\n"
        "- For reflections, if the user did not specify a title, create a concise reflective title "
        "from the description and call the function.\n"
        "- SPECIAL CASE: '谈话记录(Conversation)' is NOT a club. It is the student's overall "
        "end-of-semester CAS reflection across Creativity, Activity, and Service. If the user asks "
        "for a 'Conversation' reflection, treat it as their holistic semester CAS reflection — never "
        "describe it as a 'Conversation Club' or invent activities for such a club.\n"
        "- Valid learning outcomes: Awareness, Challenge, Initiative, Collaboration, Commitment, "
        "Global Value, Ethics, New Skills.\n\n"
    )
    if clubs:
        system_content += "Available clubs (choose exactly one of these names):\n- " + "\n- ".join(clubs) + "\n\n"
    else:
        system_content += "No clubs have been loaded yet; if the user wants to fill a form, ask them to log in / fetch clubs first.\n\n"
    if CAS_KNOWLEDGE_BASE:
        system_content += "以下是学校官方提供的IB CAS文档内容，请以此为依据回答问题：\n\n" + CAS_KNOWLEDGE_BASE

    messages = [{"role": "system", "content": system_content}]
    if attachments:
        blocks = "\n\n".join(f"[Attached file: {nm}]\n{tx}" for nm, tx in attachments)
        messages.append({"role": "system", "content":
                         "The user attached the following file(s). Use their content as context when "
                         "answering or filling forms:\n\n" + blocks})
    fixed_chars = sum(len(m.get("content", "")) for m in messages) + len(user_message)
    history_budget = max(0, CHAT_CONTEXT_CHAR_BUDGET - fixed_chars)
    messages.extend(_chat_history_messages_for_budget(history, history_budget))
    messages.append({"role": "user", "content": user_message})

    try:
        # Keep tool routing in non-thinking mode. Qwen mixed-thinking models can
        # otherwise spend the turn in reasoning_content and fail to emit a function
        # call, which leaves the user without an approval card.
        resp = llm_chat(LLM_API_KEY, LLM_MODEL, messages,
                             temperature=0.3, max_tokens=800, tools=CAS_TOOLS, tool_choice="auto",
                             thinking_enabled=False)
        msg = resp["choices"][0]["message"]
        text = (msg.get("content") or "").strip()
        reasoning = (msg.get("reasoning_content") or "").strip()
        call = _first_tool_call(msg)

        # Fallback: the model said it would act (or returned nothing) but never
        # emitted a tool call. Force the call so the user actually gets a card.
        if not call and clubs:
            stalled = bool(_INTENT_EN.search(text) or _INTENT_ZH.search(text))
            if stalled or not text:
                forced = _force_tool_call(messages, text, thinking_enabled=False)
                if forced:
                    call = forced

        if not call and thinking_enabled:
            thinking_resp = llm_chat(LLM_API_KEY, LLM_MODEL, messages,
                                     temperature=0.3, max_tokens=800,
                                     thinking_enabled=True)
            thinking_msg = thinking_resp["choices"][0]["message"]
            text = (thinking_msg.get("content") or "").strip()
            reasoning = (thinking_msg.get("reasoning_content") or "").strip()

        if call:
            fn = call.get("function", {}) or {}
            name = fn.get("name", "")
            builder = _PROPOSAL_BUILDERS.get(name)
            if not builder:
                return jsonify({"reply": text or "Sorry, I can't perform that action."})
            try:
                args = json.loads(fn.get("arguments", "{}") or "{}")
            except Exception:
                args = {}
            missing = _missing_required_details(name, args, history, user_message, attachments, clubs)
            if missing:
                return jsonify({"reply": _missing_followup(name, missing, user_message)})
            try:
                proposal = builder(args)
            except ValueError as ve:
                # Missing/invalid params: turn into a friendly follow-up question.
                return jsonify({"reply": text or f"I need a bit more info before I can do that: {ve}"})
            return jsonify({"proposal": proposal, "reply": text, "reasoning": reasoning})

        return jsonify({"reply": text, "reasoning": reasoning})
    except Exception as e:
        return jsonify({"error": sanitize_error(e)}), 500


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    # Default to 127.0.0.1 for local development; set HOST=0.0.0.0 for Docker/production
    host = os.environ.get("HOST", "127.0.0.1")
    # NOTE: For production, use a proper WSGI server (gunicorn + eventlet) instead of Werkzeug
    socketio.run(app, host=host, port=port, debug=False, allow_unsafe_werkzeug=True)
